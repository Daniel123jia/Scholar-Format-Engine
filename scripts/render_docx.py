from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from common import load_json, load_yaml, normalize_hex, format_template

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def set_run_font(run, font_cfg: Dict[str, Any], bold_override=None, italic_override=None, code=False):
    latin = "Consolas" if code else font_cfg.get("latin", "Times New Roman")
    east_asia = font_cfg.get("east_asia", latin)
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:cs"), latin)
    run.font.size = Pt(float(font_cfg.get("size_pt", 11)))
    run.bold = font_cfg.get("bold", False) if bold_override is None else bold_override
    run.italic = font_cfg.get("italic", False) if italic_override is None else italic_override
    color = normalize_hex(font_cfg.get("color", "000000"))
    run.font.color.rgb = RGBColor.from_string(color)


def apply_paragraph_format(paragraph, role_cfg: Dict[str, Any]):
    p_cfg = role_cfg.get("paragraph", {})
    paragraph.alignment = ALIGN.get(p_cfg.get("alignment", "left"), WD_ALIGN_PARAGRAPH.LEFT)
    pf = paragraph.paragraph_format
    pf.line_spacing = float(p_cfg.get("line_spacing", 1.15))
    pf.space_before = Pt(float(p_cfg.get("space_before_pt", 0)))
    pf.space_after = Pt(float(p_cfg.get("space_after_pt", 0)))
    pf.left_indent = Cm(float(p_cfg.get("left_indent_cm", 0)))
    pf.right_indent = Cm(float(p_cfg.get("right_indent_cm", 0)))
    chars = float(p_cfg.get("first_line_indent_chars", 0))
    if chars:
        font_size = float(role_cfg.get("font", {}).get("size_pt", 11))
        pf.first_line_indent = Pt(font_size * chars)
    else:
        pf.first_line_indent = Pt(0)
    pf.keep_with_next = bool(p_cfg.get("keep_with_next", False))


def get_role_cfg(style: Dict[str, Any], role: str) -> Dict[str, Any]:
    roles = style.get("roles", {})
    return roles.get(role) or roles.get("body") or next(iter(roles.values()))


def ensure_custom_style(doc: Document, name: str, role_cfg: Dict[str, Any]):
    styles = doc.styles
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    apply_paragraph_format(st.paragraph_format._parent if False else _StyleParagraphProxy(st), role_cfg)
    font_cfg = role_cfg.get("font", {})
    st.font.name = font_cfg.get("latin", "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:ascii"), font_cfg.get("latin", "Times New Roman"))
    st._element.rPr.rFonts.set(qn("w:hAnsi"), font_cfg.get("latin", "Times New Roman"))
    st._element.rPr.rFonts.set(qn("w:eastAsia"), font_cfg.get("east_asia", font_cfg.get("latin", "Times New Roman")))
    st.font.size = Pt(float(font_cfg.get("size_pt", 11)))
    st.font.bold = bool(font_cfg.get("bold", False))
    st.font.italic = bool(font_cfg.get("italic", False))
    st.font.color.rgb = RGBColor.from_string(normalize_hex(font_cfg.get("color", "000000")))
    return st


class _StyleParagraphProxy:
    """Tiny adapter exposing paragraph_format-like attributes for a style."""
    def __init__(self, style):
        self.paragraph_format = style.paragraph_format
        self.alignment = None


def configure_style(doc: Document, style_name: str, role_cfg: Dict[str, Any]):
    try:
        st = doc.styles[style_name]
    except KeyError:
        st = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    font_cfg = role_cfg.get("font", {})
    st.font.name = font_cfg.get("latin", "Times New Roman")
    rpr = st._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), font_cfg.get("latin", "Times New Roman"))
    rfonts.set(qn("w:hAnsi"), font_cfg.get("latin", "Times New Roman"))
    rfonts.set(qn("w:eastAsia"), font_cfg.get("east_asia", font_cfg.get("latin", "Times New Roman")))
    st.font.size = Pt(float(font_cfg.get("size_pt", 11)))
    st.font.bold = bool(font_cfg.get("bold", False))
    st.font.italic = bool(font_cfg.get("italic", False))
    st.font.color.rgb = RGBColor.from_string(normalize_hex(font_cfg.get("color", "000000")))

    p_cfg = role_cfg.get("paragraph", {})
    pf = st.paragraph_format
    pf.alignment = ALIGN.get(p_cfg.get("alignment", "left"), WD_ALIGN_PARAGRAPH.LEFT)
    pf.line_spacing = float(p_cfg.get("line_spacing", 1.15))
    pf.space_before = Pt(float(p_cfg.get("space_before_pt", 0)))
    pf.space_after = Pt(float(p_cfg.get("space_after_pt", 0)))
    pf.left_indent = Cm(float(p_cfg.get("left_indent_cm", 0)))
    pf.right_indent = Cm(float(p_cfg.get("right_indent_cm", 0)))
    chars = float(p_cfg.get("first_line_indent_chars", 0))
    font_size = float(font_cfg.get("size_pt", 11))
    pf.first_line_indent = Pt(font_size * chars) if chars else Pt(0)
    pf.keep_with_next = bool(p_cfg.get("keep_with_next", False))
    return st


def setup_styles(doc: Document, style: Dict[str, Any]):
    roles = style.get("roles", {})
    mapping = {
        "title": "Title",
        "heading1": "Heading 1",
        "heading2": "Heading 2",
        "heading3": "Heading 3",
        "body": "Normal",
    }
    for role, style_name in mapping.items():
        if role in roles:
            configure_style(doc, style_name, roles[role])
    for role, cfg in roles.items():
        if role in mapping:
            continue
        configure_style(doc, f"SDD {role}", cfg)


def set_page_layout(doc: Document, style: Dict[str, Any]):
    cfg = style.get("page", {})
    margins = cfg.get("margins_mm", {})
    for section in doc.sections:
        if cfg.get("size") == "LETTER":
            section.page_width = Mm(215.9)
            section.page_height = Mm(279.4)
        else:
            section.page_width = Mm(210)
            section.page_height = Mm(297)
        if cfg.get("orientation") == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Mm(float(margins.get("top", 25.4)))
        section.bottom_margin = Mm(float(margins.get("bottom", 25.4)))
        section.left_margin = Mm(float(margins.get("left", 25.4)))
        section.right_margin = Mm(float(margins.get("right", 25.4)))


def add_field(run, instr: str, display_text: str = ""):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display_text
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2, text, fld_char3])


def setup_header_footer(doc: Document, document: Dict[str, Any], style: Dict[str, Any]):
    features = style.get("features", {})
    context = {"title": document.get("title") or "", **document.get("metadata", {})}
    header_text = format_template(features.get("header_text"), context)
    footer_text = format_template(features.get("footer_text"), context)
    for section in doc.sections:
        if header_text:
            p = section.header.paragraphs[0]
            p.text = header_text
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cfg = get_role_cfg(style, "note")
            apply_paragraph_format(p, cfg)
            for run in p.runs:
                set_run_font(run, cfg.get("font", {}))
        if footer_text or features.get("page_numbers"):
            p = section.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cfg = get_role_cfg(style, "note")
            apply_paragraph_format(p, cfg)
            if footer_text:
                r = p.add_run(footer_text)
                set_run_font(r, cfg.get("font", {}))
                if features.get("page_numbers"):
                    r2 = p.add_run(" · ")
                    set_run_font(r2, cfg.get("font", {}))
            if features.get("page_numbers"):
                r = p.add_run()
                set_run_font(r, cfg.get("font", {}))
                add_field(r, "PAGE", "1")


def add_hyperlink(paragraph, text: str, url: str, role_cfg: Dict[str, Any]):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(underline)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def style_paragraph(paragraph, role: str, style: Dict[str, Any]):
    cfg = get_role_cfg(style, role)
    apply_paragraph_format(paragraph, cfg)
    for run in paragraph.runs:
        set_run_font(run, cfg.get("font", {}))


def add_rich_paragraph(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    role = block.get("role", "body")
    style_name = "Normal" if role == "body" else f"SDD {role}"
    if role in {"heading1", "heading2", "heading3"}:
        style_name = {"heading1": "Heading 1", "heading2": "Heading 2", "heading3": "Heading 3"}[role]
    try:
        p = doc.add_paragraph(style=style_name)
    except Exception:
        p = doc.add_paragraph()
    cfg = get_role_cfg(style, role)
    apply_paragraph_format(p, cfg)
    if block.get("text") is not None:
        r = p.add_run(str(block.get("text") or ""))
        set_run_font(r, cfg.get("font", {}))
    else:
        for item in block.get("runs", []):
            if item.get("link"):
                add_hyperlink(p, item.get("text", ""), item["link"], cfg)
                continue
            r = p.add_run(item.get("text", ""))
            set_run_font(
                r,
                cfg.get("font", {}),
                bold_override=item.get("bold", cfg.get("font", {}).get("bold", False)),
                italic_override=item.get("italic", cfg.get("font", {}).get("italic", False)),
                code=item.get("code", False),
            )
            r.underline = bool(item.get("underline", False))
    return p


def add_heading(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    level = min(max(int(block.get("level", 1)), 1), 6)
    style_name = f"Heading {level}"
    if level > 3:
        try:
            base = get_role_cfg(style, "heading3")
            configure_style(doc, style_name, base)
        except Exception:
            pass
    p = doc.add_paragraph(style=style_name)
    role = f"heading{min(level, 3)}"
    cfg = get_role_cfg(style, role)
    apply_paragraph_format(p, cfg)
    r = p.add_run(block.get("text", ""))
    set_run_font(r, cfg.get("font", {}))
    return p


def set_cell_shading(cell, fill: str | None):
    if not fill:
        return
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), normalize_hex(fill, "FFFFFF"))


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = "w:" + edge
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, value in edge_data.items():
                element.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def style_table(table, style: Dict[str, Any]):
    cfg = style.get("table", {})
    table.autofit = bool(cfg.get("autofit", True))
    preset = cfg.get("preset", "three_line")
    border_color = normalize_hex(cfg.get("border_color", "808080"), "808080")
    all_cells = [cell for row in table.rows for cell in row.cells]

    if preset == "grid":
        border = {"val": "single", "sz": "6", "color": border_color}
        for cell in all_cells:
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    elif preset == "minimal":
        border = {"val": "single", "sz": "4", "color": "D1D5DB"}
        for cell in all_cells:
            set_cell_border(cell, bottom=border)
    else:
        none = {"val": "nil"}
        for cell in all_cells:
            set_cell_border(cell, top=none, bottom=none, left=none, right=none)
        if table.rows:
            thick = {"val": "single", "sz": "10", "color": border_color}
            thin = {"val": "single", "sz": "6", "color": border_color}
            for cell in table.rows[0].cells:
                set_cell_border(cell, top=thick, bottom=thin)
            for cell in table.rows[-1].cells:
                set_cell_border(cell, bottom=thick)


def add_table(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    if block.get("caption"):
        p = doc.add_paragraph(style="SDD caption")
        cfg = get_role_cfg(style, "caption")
        apply_paragraph_format(p, cfg)
        r = p.add_run(block["caption"])
        set_run_font(r, cfg.get("font", {}))

    headers = block.get("headers", [])
    rows = block.get("rows", [])
    ncols = len(headers) or max([len(r) for r in rows], default=1)
    nrows = len(rows) + (1 if headers else 0)
    table = doc.add_table(rows=max(nrows, 1), cols=max(ncols, 1))
    table.alignment = 1

    row_offset = 0
    if headers:
        row = table.rows[0]
        row_offset = 1
        if style.get("table", {}).get("repeat_header"):
            set_repeat_table_header(row)
        fill = style.get("table", {}).get("header_fill")
        for j, value in enumerate(headers):
            cell = row.cells[j]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, fill)
            p = cell.paragraphs[0]
            cfg = get_role_cfg(style, "table_header")
            apply_paragraph_format(p, cfg)
            r = p.add_run(str(value))
            set_run_font(r, cfg.get("font", {}))

    for i, row_values in enumerate(rows):
        row = table.rows[i + row_offset]
        for j in range(ncols):
            value = row_values[j] if j < len(row_values) else ""
            cell = row.cells[j]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            cfg = get_role_cfg(style, "table")
            apply_paragraph_format(p, cfg)
            r = p.add_run("" if value is None else str(value))
            set_run_font(r, cfg.get("font", {}))

    style_table(table, style)

    if block.get("note"):
        p = doc.add_paragraph(style="SDD note")
        cfg = get_role_cfg(style, "note")
        apply_paragraph_format(p, cfg)
        r = p.add_run("Note: " + str(block["note"]))
        set_run_font(r, cfg.get("font", {}))
    return table


def add_callout(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    role = block.get("role", "note")
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    cell.text = ""
    call_cfg = (style.get("callouts") or {}).get(role, {})
    set_cell_shading(cell, call_cfg.get("fill"))
    border_color = normalize_hex(call_cfg.get("border_color"), "9CA3AF")
    set_cell_border(
        cell,
        left={"val": "single", "sz": "18", "color": border_color},
        top={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"}
    )
    p = cell.paragraphs[0]
    cfg = get_role_cfg(style, role)
    apply_paragraph_format(p, cfg)
    if block.get("title"):
        r = p.add_run(str(block["title"]) + "\n")
        set_run_font(r, cfg.get("font", {}), bold_override=True)
    r = p.add_run(str(block.get("text", "")))
    set_run_font(r, cfg.get("font", {}))
    meta = block.get("meta") or {}
    if meta:
        r = p.add_run("\n")
        set_run_font(r, cfg.get("font", {}))
        for idx, (k, v) in enumerate(meta.items()):
            if v is None or not str(v).strip():
                continue
            if idx:
                r = p.add_run(" · ")
                set_run_font(r, cfg.get("font", {}))
            r = p.add_run(f"{k}: {v}")
            set_run_font(r, cfg.get("font", {}), italic_override=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_quote(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    p = doc.add_paragraph(style="SDD quote")
    cfg = get_role_cfg(style, "quote")
    apply_paragraph_format(p, cfg)
    r = p.add_run(block.get("text", ""))
    set_run_font(r, cfg.get("font", {}))
    if block.get("source"):
        r = p.add_run("\n— " + str(block["source"]))
        set_run_font(r, cfg.get("font", {}), italic_override=True)
    return p


def add_list(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    role = block.get("role", "body")
    cfg = get_role_cfg(style, role)
    for item in block.get("items", []):
        style_name = "List Number" if block.get("ordered") else "List Bullet"
        p = doc.add_paragraph(style=style_name)
        apply_paragraph_format(p, cfg)
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(str(item))
        set_run_font(r, cfg.get("font", {}))


def add_image(doc: Document, block: Dict[str, Any], style: Dict[str, Any], input_base: Path):
    path = Path(block.get("path", ""))
    if not path.is_absolute():
        path = input_base / path
    if not path.exists():
        p = doc.add_paragraph(style="SDD warning")
        cfg = get_role_cfg(style, "warning")
        apply_paragraph_format(p, cfg)
        r = p.add_run(f"[Missing image: {path}]")
        set_run_font(r, cfg.get("font", {}))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width = block.get("width_cm")
    if width:
        p.add_run().add_picture(str(path), width=Cm(float(width)))
    else:
        p.add_run().add_picture(str(path), width=Cm(14.5))
    if block.get("caption"):
        cp = doc.add_paragraph(style="SDD caption")
        cfg = get_role_cfg(style, "caption")
        apply_paragraph_format(cp, cfg)
        r = cp.add_run(str(block["caption"]))
        set_run_font(r, cfg.get("font", {}))


def add_equation(doc: Document, block: Dict[str, Any], style: Dict[str, Any]):
    p = doc.add_paragraph(style="SDD equation")
    cfg = get_role_cfg(style, "equation")
    apply_paragraph_format(p, cfg)
    text = str(block.get("text", ""))
    if block.get("label"):
        text = f"{text}    {block['label']}"
    r = p.add_run(text)
    set_run_font(r, cfg.get("font", {}))


def add_horizontal_rule(doc: Document):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D1D5DB")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_toc(doc: Document, style: Dict[str, Any], document: Dict[str, Any]):
    title = style.get("features", {}).get("toc_title", "Contents")
    p = doc.add_paragraph()
    cfg = get_role_cfg(style, "heading1")
    apply_paragraph_format(p, cfg)
    r = p.add_run(title)
    set_run_font(r, cfg.get("font", {}))

    # v1.0-alpha uses a deterministic static TOC so previews and exports
    # remain stable across Word/LibreOffice without requiring field refresh.
    body_cfg = get_role_cfg(style, "body")
    for block in document.get("blocks", []):
        if block.get("type") == "heading" and int(block.get("level", 1)) == 1:
            tp = doc.add_paragraph()
            apply_paragraph_format(tp, body_cfg)
            tp.paragraph_format.first_line_indent = Pt(0)
            tp.paragraph_format.left_indent = Cm(0.4)
            tr = tp.add_run(str(block.get("text", "")))
            set_run_font(tr, body_cfg.get("font", {}))


def render(document: Dict[str, Any], style: Dict[str, Any], output: str | Path, input_base: str | Path | None = None):
    doc = Document()
    setup_styles(doc, style)
    set_page_layout(doc, style)
    input_base = Path(input_base or ".")

    # Title
    if document.get("title"):
        p = doc.add_paragraph(style="Title")
        cfg = get_role_cfg(style, "title")
        apply_paragraph_format(p, cfg)
        r = p.add_run(str(document["title"]))
        set_run_font(r, cfg.get("font", {}))
    if document.get("subtitle"):
        p = doc.add_paragraph(style="SDD subtitle")
        cfg = get_role_cfg(style, "subtitle")
        apply_paragraph_format(p, cfg)
        r = p.add_run(str(document["subtitle"]))
        set_run_font(r, cfg.get("font", {}))

    # Metadata
    if style.get("features", {}).get("metadata_table") and document.get("metadata"):
        rows = [[k, v] for k, v in document["metadata"].items() if v is not None and str(v).strip()]
        if rows:
            add_table(doc, {"headers": ["项目", "内容"], "rows": rows, "role": "table"}, style)
            doc.add_paragraph()

    toc_inserted = False
    for block in document.get("blocks", []):
        if style.get("features", {}).get("toc") and not toc_inserted and block.get("type") == "heading" and int(block.get("level", 1)) == 1:
            add_toc(doc, style, document)
            if style.get("features", {}).get("toc_page_break_after", True):
                doc.add_page_break()
            toc_inserted = True

        t = block.get("type")
        if t == "heading":
            add_heading(doc, block, style)
        elif t == "paragraph":
            add_rich_paragraph(doc, block, style)
        elif t == "list":
            add_list(doc, block, style)
        elif t == "table":
            add_table(doc, block, style)
        elif t == "quote":
            add_quote(doc, block, style)
        elif t == "callout":
            add_callout(doc, block, style)
        elif t == "image":
            add_image(doc, block, style, input_base)
        elif t == "equation":
            add_equation(doc, block, style)
        elif t == "page_break":
            doc.add_page_break()
        elif t == "horizontal_rule":
            add_horizontal_rule(doc)

    setup_header_footer(doc, document, style)

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    # Re-open as a basic integrity check.
    Document(str(output))
    return output


def main() -> int:
    ap = argparse.ArgumentParser(description="Render Scholar DocumentIR to DOCX.")
    ap.add_argument("--input", required=True)
    ap.add_argument("--style", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    document = load_json(args.input)
    style = load_yaml(args.style)
    render(document, style, args.output, Path(args.input).resolve().parent)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
